#!/usr/bin/env python3
# =============================================================================
# excel_to_okf.py — Knowledge Plane workbook (xlsx) -> OKF bundle (kp/)
#
#   Business authors + approves in Excel; this writes the canonical bundle.
#   Safe to re-run: preserves the exporter's GENERATED blocks, build-stamped
#   `last_validated`, and non-namespace keys in bundle.yaml.
#   RULE: when Excel is the authoring surface, it is the SOLE write path for
#   the columns it owns — do not also hand-edit those frontmatter fields.
#
#   Usage: python3 excel_to_okf.py [workbook.xlsx] [out_dir]
# =============================================================================
import sys, os, re, datetime
import openpyxl
from okf_slug import slugify  # OKF-8: one shared slug, agrees with okf-slug.js

WB_FILE = sys.argv[1] if len(sys.argv) > 1 else 'knowledge_plane_workbook.xlsx'
OUT     = sys.argv[2] if len(sys.argv) > 2 else 'kp'

KINDS = {'entity', 'defined_class', 'measure', 'attribute'}
# OKF v0.2 lifecycle (SPEC.md §5.4): draft | stable | deprecated. 'In review'
# has no v0.2 equivalent, so it falls back to 'draft' via STATUS_MAP.get(...).
STATUS_MAP = {'Approved': 'stable', 'Draft': 'draft', 'Deprecated': 'deprecated'}
STEWARD_TO_DOMAIN = {'Sales': 'sales', 'Finance': 'finance',
                     'Merchandising': 'merchandising', 'Retail Operations': 'operations'}
GEN_BEGIN = '<!-- BEGIN GENERATED: implementations (written by exporter; do not edit) -->'
GEN_END   = '<!-- END GENERATED -->'

errors, warnings = [], []
wb = openpyxl.load_workbook(WB_FILE, data_only=True)

def sheet_rows(name):
    ws = wb[name]
    headers = [str(c.value).strip() if c.value else '' for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        row = {h: (v.strip() if isinstance(v, str) else v) for h, v in zip(headers, r) if h}
        if row.get('URI'):
            yield row

concepts, order = {}, []
for row in sheet_rows('Concepts'):
    uri = row['URI']
    if uri in concepts: errors.append(f'duplicate URI on Concepts sheet: {uri}'); continue
    concepts[uri] = row; order.append(uri)
relationships = list(sheet_rows('Relationships'))

namespace = None
for r in wb['Read me'].iter_rows(values_only=True):
    for v in r:
        if isinstance(v, str) and 'Namespace:' in v:
            m = re.search(r'Namespace:\s*(\S+)', v)
            if m: namespace = m.group(1)
if not namespace: errors.append('namespace not found on the Read me sheet ("Namespace: <uri>")')

# ---- validate ----------------------------------------------------------------
for uri, c in concepts.items():
    kind = c.get('Kind')
    if kind not in KINDS: errors.append(f'{uri}: invalid Kind "{kind}"')
    if not c.get('Label'):      errors.append(f'{uri}: missing Label')
    if not c.get('Definition'): errors.append(f'{uri}: missing Definition')
    st = c.get('Approval status')
    if st not in STATUS_MAP: warnings.append(f'{uri}: Approval status "{st}" -> treated as Draft')
    if kind == 'defined_class':
        if not c.get('Membership rule'): errors.append(f'{uri}: defined_class requires a Membership rule')
        if not c.get('Subtype of'):      errors.append(f'{uri}: defined_class requires Subtype of')
    elif c.get('Membership rule'):
        warnings.append(f'{uri}: Membership rule set but Kind is {kind}')
    for col in ('Subtype of', 'Of (entity)'):
        v = c.get(col)
        if v and v not in concepts: errors.append(f'{uri}: {col} "{v}" is not on the Concepts sheet')
    of = c.get('Of (entity)')
    if of and of in concepts and concepts[of].get('Kind') != 'entity':
        errors.append(f'{uri}: Of (entity) {of} must be an entity')
    if st == 'Approved' and not c.get('Approved by'):
        warnings.append(f'{uri}: Approved but "Approved by" is empty')

rel_by_domain = {}
for r in relationships:
    dom, rng = r.get('From (domain)'), r.get('To (range)')
    if dom not in concepts: errors.append(f'relationship {r["URI"]}: domain "{dom}" not on Concepts sheet')
    elif concepts[dom].get('Kind') != 'entity': errors.append(f'relationship {r["URI"]}: domain must be an entity')
    if rng not in concepts: errors.append(f'relationship {r["URI"]}: range "{rng}" not on Concepts sheet')
    if not r.get('Verb phrase'): errors.append(f'relationship {r["URI"]}: missing Verb phrase')
    if dom: rel_by_domain.setdefault(dom, []).append(r)

if errors:
    print('GENERATION FAILED — workbook validation:')
    for e in errors: print('  ERROR ' + e)
    sys.exit(1)
for w in warnings: print('  WARN ' + w)

# ---- emit --------------------------------------------------------------------
def domain_of(c): return STEWARD_TO_DOMAIN.get(c.get('Steward'), 'global')
def kebab(uri):   return re.sub(r'(?<!^)(?=[A-Z])', '-', uri.split(':')[1]).lower()
path_of = {u: f'{domain_of(c)}/{kebab(u)}.md' for u, c in concepts.items()}

def yq(s):
    if isinstance(s, list): return '[' + ', '.join(yq(x) for x in s) + ']'
    s = str(s)
    return s if re.match(r'^[A-Za-z0-9_\-:.]+$', s) else '"' + s.replace('"', '\\"') + '"'

def rel_link(from_uri, to_uri):
    return os.path.relpath(path_of[to_uri], os.path.dirname(path_of[from_uri])).replace(os.sep, '/')

def preserved(fp):
    """Keep the build's write-backs across regeneration: GENERATED block + last_validated."""
    gen, lv = None, None
    if os.path.exists(fp):
        t = open(fp, encoding='utf-8').read()
        b, e = t.find(GEN_BEGIN), t.find(GEN_END)
        if b != -1 and e > b: gen = t[b + len(GEN_BEGIN):e].strip('\n')
        m = re.search(r'^last_validated:\s*(.+)$', t, re.M)
        if m and m.group(1).strip() not in ('null', ''): lv = m.group(1).strip()
    return gen, lv

today = datetime.date.today().isoformat()
os.makedirs(OUT, exist_ok=True)
counts = {}
for uri in order:
    c = concepts[uri]
    p = os.path.join(OUT, path_of[uri])
    os.makedirs(os.path.dirname(p), exist_ok=True)
    gen_block, last_validated = preserved(p)
    status = STATUS_MAP.get(c.get('Approval status'), 'draft')
    ao = c.get('Approved on')
    if isinstance(ao, datetime.datetime): ao = ao.date().isoformat()
    elif isinstance(ao, datetime.date):   ao = ao.isoformat()
    syns  = [s.strip() for s in str(c.get('Synonyms') or '').split(',') if s.strip()]
    roles = [s.strip() for s in str(c.get('Allowed roles') or '').split(',') if s.strip()]

    fm = [f'uri: {yq(uri)}', f'type: {c["Kind"]}', f'title: {yq(c["Label"])}',
          f'description: {yq(c["Definition"])}']
    if syns:                      fm.append(f'synonyms: {yq(syns)}')
    if c.get('Steward'):          fm.append(f'steward: {yq(c["Steward"])}')
    if c.get('Subtype of'):       fm.append(f'subtype_of: {yq(c["Subtype of"])}')
    if c.get('Of (entity)') and c['Kind'] != 'entity':
        fm.append(f'of: {yq(c["Of (entity)"])}')
    if c.get('Membership rule'):  fm.append(f'membership_rule: {yq(c["Membership rule"])}')
    if c.get('Preferred source'): fm.append(f'preferred_source: {yq(c["Preferred source"])}')
    if roles:                     fm.append(f'allowed_roles: {yq(roles)}')
    myrels = rel_by_domain.get(uri, [])
    if myrels:
        fm.append('relationships:')
        for r in sorted(myrels, key=lambda x: x['URI']):
            fm += [f'  - uri: {yq(r["URI"])}', f'    verb: {yq(r["Verb phrase"])}',
                   f'    range: {yq(r["To (range)"])}']
    fm.append(f'tags: {yq([c.get("Steward") or "global"])}')
    fm.append(f'status: {status}')
    vs = c.get('Validation status')
    if vs and vs != 'Not started':
        fm.append(f'validation_status: {yq(str(vs).lower())}')
        if c.get('Checked against'):  fm.append(f'checked_against: {yq(c["Checked against"])}')
        if c.get('Validation notes'): fm.append(f'validation_notes: {yq(c["Validation notes"])}')
    if last_validated:            fm.append(f'last_validated: {last_validated}')
    actor = f'human:{slugify(c.get("Steward") or "global")}'
    at = f'{ao or today}T00:00:00Z'
    fm.append(f'generated: {{ by: {actor}, at: {at} }}')
    if status == 'stable':        fm.append(f'verified: {{ by: {actor}, at: {at} }}')

    body = [f'# {c["Label"]}', '', c['Definition'], '']
    if c.get('Membership rule'):
        body += ['## Membership rule', '', f'`{c["Membership rule"]}`', '',
                 'Apply this rule **verbatim** when querying; never improvise an equivalent filter.', '']
    if c.get('Subtype of'):
        t = c['Subtype of']
        body += [f'Subtype of [{concepts[t]["Label"]}]({rel_link(uri, t)}).', '']
    if c.get('Of (entity)') and c['Kind'] != 'entity':
        t = c['Of (entity)']
        noun = 'Attribute' if c['Kind'] == 'attribute' else 'Measure'
        body += [f'{noun} of [{concepts[t]["Label"]}]({rel_link(uri, t)}).', '']
    if myrels:
        body += ['## Relationships', '']
        for r in sorted(myrels, key=lambda x: x['URI']):
            tgt = r['To (range)']
            body.append(f'- *{r["Verb phrase"]}* → [{concepts[tgt]["Label"]}]({rel_link(uri, tgt)}) (`{r["URI"]}`)')
        body.append('')
    body += ['## Implementations', '', GEN_BEGIN,
             gen_block if gen_block is not None else '_Not yet generated — run `node build.js`._',
             GEN_END, '']

    with open(p, 'w', encoding='utf-8') as f:
        f.write('---\n' + '\n'.join(fm) + '\n---\n\n' + '\n'.join(body))
    counts[domain_of(c)] = counts.get(domain_of(c), 0) + 1

cfg_path = os.path.join(OUT, 'bundle.yaml')
extra = []
if os.path.exists(cfg_path):
    extra = [l.rstrip('\n') for l in open(cfg_path) if l.strip() and not l.startswith('namespace:')]
with open(cfg_path, 'w') as f:
    f.write(f'namespace: "{namespace}"\n' + ('\n'.join(extra) + '\n' if extra else ''))

# prune concepts whose URI vanished (never touches agent/, _templates/, indexes)
emitted = {os.path.normpath(os.path.join(OUT, p)) for p in path_of.values()}
for root, dirs, files in os.walk(OUT):
    dirs[:] = [d for d in dirs if not d.startswith('_') and d != 'agent']
    for fn in files:
        fp = os.path.normpath(os.path.join(root, fn))
        if fn.endswith('.md') and fn not in ('index.md', 'README.md') and fp not in emitted:
            os.remove(fp); print(f'  pruned (no longer in workbook): {os.path.relpath(fp, OUT)}')

print(f'\nBundle written to {OUT}/: {len(concepts)} concepts '
      f'({", ".join(f"{d}: {n}" for d, n in sorted(counts.items()))}), {len(relationships)} relationships')
print('Next: node build.js')
